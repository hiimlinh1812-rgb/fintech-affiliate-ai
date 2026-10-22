"""
Reporting & Export Service
Generates Excel (.xlsx) and PDF (.pdf) reports reflecting the current filtered view.
Complies with Section 11 of requirements and Section 1.3.6 of CNDVTC.docx.
"""

import io
from datetime import datetime
from typing import Dict, Any, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors


def export_to_excel(df: pd.DataFrame, platform_summary: pd.DataFrame, month_summary: pd.DataFrame, filters: Dict[str, str]) -> bytes:
    """
    Creates a multi-tab professional Excel workbook:
      - Sheet 1: Tổng quan KPI & Lọc
      - Sheet 2: Báo cáo theo Nền Tảng
      - Sheet 3: Báo cáo theo Tháng
      - Sheet 4: Chi tiết Dữ liệu
    """
    wb = openpyxl.Workbook()
    # Sheet 1: Overview
    ws_meta = wb.active
    ws_meta.title = "Tổng Quan & Bộ Lọc"

    # Header styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Segoe UI", size=15, bold=True, color="1E3A8A")
    regular_font = Font(name="Segoe UI", size=10)
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Title
    ws_meta["B2"] = "BÁO CÁO HIỆU QUẢ TIẾP THỊ LIÊN KẾT (FINTECH AFFILIATE DASHBOARD)"
    ws_meta["B2"].font = title_font
    ws_meta["B3"] = f"Thời gian trích xuất: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Trạng thái: Demo Học Thuật"
    ws_meta["B3"].font = regular_font

    # Filters block
    ws_meta["B5"] = "BỘ LỌC ĐANG ÁP DỤNG"
    ws_meta["B5"].font = bold_font
    ws_meta["B6"] = "Nền tảng (Platform):"
    ws_meta["C6"] = filters.get("platform", "Tất cả")
    ws_meta["B7"] = "Tháng (Month):"
    ws_meta["C7"] = filters.get("month", "Tất cả")
    ws_meta["B8"] = "Khu vực (Region):"
    ws_meta["C8"] = filters.get("region", "Tất cả")

    # KPI summary block
    total_rev = df["revenue"].sum() if not df.empty else 0
    total_comm = df["commission"].sum() if not df.empty else 0
    total_profit = df["profit"].sum() if not df.empty else 0
    total_ads = df["ad_cost"].sum() if not df.empty else 0
    avg_roi = round(total_comm / total_ads, 2) if total_ads > 0 else 0.0

    ws_meta["B10"] = "CHỈ SỐ TÀI CHÍNH CHÍNH"
    ws_meta["B10"].font = bold_font
    ws_meta["B11"] = "Tổng Doanh Thu (GMV):"
    ws_meta["C11"] = f"{int(total_rev):,} VND"
    ws_meta["B12"] = "Tổng Hoa Hồng:"
    ws_meta["C12"] = f"{int(total_comm):,} VND"
    ws_meta["B13"] = "Tổng Chi Phí Quảng Cáo:"
    ws_meta["C13"] = f"{int(total_ads):,} VND"
    ws_meta["B14"] = "Tổng Lợi Nhuận Ròng:"
    ws_meta["C14"] = f"{int(total_profit):,} VND"
    ws_meta["B15"] = "Hệ Số ROI Trung Bình:"
    ws_meta["C15"] = f"{avg_roi:.2f}x"

    for r in range(5, 16):
        ws_meta[f"B{r}"].font = bold_font
        ws_meta[f"C{r}"].font = regular_font

    ws_meta.column_dimensions["B"].width = 28
    ws_meta.column_dimensions["C"].width = 30

    # Sheet 2: Platform Breakdown
    if not platform_summary.empty:
        ws_plat = wb.create_sheet(title="Theo Nền Tảng")
        cols = list(platform_summary.columns)
        ws_plat.append(cols)
        for _, row in platform_summary.iterrows():
            ws_plat.append(list(row))
        _style_sheet_headers(ws_plat, header_fill, header_font, border_thin)

    # Sheet 3: Monthly Breakdown
    if not month_summary.empty:
        ws_month = wb.create_sheet(title="Theo Tháng")
        cols = list(month_summary.columns)
        ws_month.append(cols)
        for _, row in month_summary.iterrows():
            ws_month.append(list(row))
        _style_sheet_headers(ws_month, header_fill, header_font, border_thin)

    # Sheet 4: Raw / Filtered Records
    ws_data = wb.create_sheet(title="Dữ Liệu Chi Tiết")
    if not df.empty:
        display_cols = [c for c in df.columns if c != "month"]
        ws_data.append(display_cols)
        for _, row in df[display_cols].iterrows():
            ws_data.append(list(row))
        _style_sheet_headers(ws_data, header_fill, header_font, border_thin)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _style_sheet_headers(ws, fill, font, border):
    """Utility to style header row and auto-adjust column width."""
    for col_idx, col in enumerate(ws.iter_cols(1, ws.max_column), 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
        max_len = max(len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(1, min(ws.max_row + 1, 30)))
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)


def export_to_pdf(df: pd.DataFrame, platform_summary: pd.DataFrame, filters: Dict[str, str]) -> bytes:
    """
    Generates a PDF report using ReportLab.
    Contains summary metadata, KPI table, Platform breakdown, and Academic Disclaimer.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Normal"],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#1E3A8A"),
        fontName="Helvetica-Bold",
        spaceAfter=6
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748B"),
        fontName="Helvetica",
        spaceAfter=14
    )
    heading_style = ParagraphStyle(
        "Heading",
        parent=styles["Normal"],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#0F172A"),
        fontName="Helvetica-Bold",
        spaceBefore=10,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#334155"),
        fontName="Helvetica"
    )

    story = []

    # Title & Subtitle
    story.append(Paragraph("FINTECH AFFILIATE DASHBOARD - BÁO CÁO DOANH THU", title_style))
    story.append(Paragraph(
        f"Xuất ngày: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | "
        f"Bộ lọc: Sàn: {filters.get('platform', 'Tất cả')} | "
        f"Tháng: {filters.get('month', 'Tất cả')} | "
        f"Khu vực: {filters.get('region', 'Tất cả')}",
        subtitle_style
    ))
    story.append(Spacer(1, 8))

    # KPI Summary Section
    story.append(Paragraph("1. Tổng Hợp Chỉ Số Hiệu Suất Chính (KPI)", heading_style))

    total_rev = df["revenue"].sum() if not df.empty else 0
    total_comm = df["commission"].sum() if not df.empty else 0
    total_profit = df["profit"].sum() if not df.empty else 0
    total_ads = df["ad_cost"].sum() if not df.empty else 0
    avg_roi = round(total_comm / total_ads, 2) if total_ads > 0 else 0.0
    total_orders = df["orders"].sum() if not df.empty else 0
    total_clicks = df["clicks"].sum() if not df.empty else 0
    cr = round(total_orders / total_clicks * 100, 2) if total_clicks > 0 else 0.0

    kpi_data = [
        ["Chỉ số", "Giá trị đạt được", "Đơn vị tính"],
        ["Tổng Doanh Thu (GMV)", f"{int(total_rev):,} VND", "Doanh thu bán hàng"],
        ["Tổng Hoa Hồng", f"{int(total_comm):,} VND", "Thu nhập sàn chi trả"],
        ["Tổng Chi Phí Quảng Cáo", f"{int(total_ads):,} VND", "Chi phí chiến dịch"],
        ["Tổng Lợi Nhuận Ròng", f"{int(total_profit):,} VND", "Hoa hồng - Chi phí"],
        ["ROI Trung Bình", f"{avg_roi:.2f}x", "Tỷ suất sinh lời"],
        ["Tỷ Lệ Chuyển Đổi (CR)", f"{cr:.2f}%", "Đơn hàng / Lượt nhấp"],
        ["Tổng Số Đơn Hàng", f"{int(total_orders):,}", "Đơn hàng thành công"]
    ]

    t_kpi = Table(kpi_data, colWidths=[180, 180, 160])
    t_kpi.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
    ]))
    story.append(t_kpi)
    story.append(Spacer(1, 14))

    # Platform Breakdown Section
    story.append(Paragraph("2. So Sánh Hiệu Quả Các Nền Tảng (Shopee, Lazada, TikTok Shop)", heading_style))

    if not platform_summary.empty:
        plat_rows = [["Nền Tảng", "Doanh Thu (GMV)", "Hoa Hồng", "Chi Phí Ads", "Lợi Nhuận", "ROI"]]
        for _, r in platform_summary.iterrows():
            plat_rows.append([
                str(r.get("platform", "")),
                f"{int(r.get('revenue', 0)):,} ₫",
                f"{int(r.get('commission', 0)):,} ₫",
                f"{int(r.get('ad_cost', 0)):,} ₫",
                f"{int(r.get('profit', 0)):,} ₫",
                f"{r.get('roi', 0):.2f}x"
            ])
        t_plat = Table(plat_rows, colWidths=[90, 95, 95, 85, 95, 60])
        t_plat.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F766E')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F0FDFA')])
        ]))
        story.append(t_plat)
    else:
        story.append(Paragraph("Không có dữ liệu cho bộ lọc hiện tại.", body_style))

    story.append(Spacer(1, 14))

    # Academic Disclaimer
    story.append(Paragraph("3. Ghi Chú và Giới Hạn Nguyên Mẫu (Academic Disclaimer)", heading_style))
    story.append(Paragraph(
        "Báo cáo này được kết xuất tự động từ hệ thống Nguyên mẫu Nghiên cứu (Prototype) "
        "dựa trên dữ liệu mô phỏng theo nội dung báo cáo CNDVTC.docx. "
        "Hệ thống không sử dụng dữ liệu thực tế và không kết nối API trực tiếp từ các sàn TMĐT. "
        "Các chỉ số lợi nhuận và dự báo mang tính chất thử nghiệm học thuật.",
        body_style
    ))

    doc.build(story)
    return buffer.getvalue()
